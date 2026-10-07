"""Read Bloomberg breakdown exports without changing or filling the workbook."""
import json
import math
from numbers import Real
from pathlib import Path

import pandas as pd
from .analysis import markdown_table, sha256
from .data import save_bundle


SCHEMAS = {
    'LBUSTRUU_Sector': ('us_aggregate', 'sector', ('Sector', 'Weight (%)')),
    'LBUSTRUU_Rating': ('us_aggregate', 'rating', ('Rating', 'Weight (%)')),
    'LG38TRUH_Currency': ('global_ex_us', 'currency', ('Currency', 'Weight (%)')),
    'LG38TRUH_Local_Yield_Duration': (
        'global_ex_us', 'currency_analytics',
        ('Currency', 'Wtd-Avg Duration (yrs)', 'Wtd-Avg Local Yield (%)')),
}
MISSING = {'', '-', '\u2014', '\u2013', 'N/A', 'NA', '#N/A', 'N.A.'}


def _number(value, location, metric):
    if value is None or (isinstance(value, str) and value.strip().upper() in MISSING):
        return float('nan')
    if isinstance(value, bool) or not isinstance(value, Real) or not math.isfinite(value):
        raise ValueError(f'Expected a finite numeric {metric} at {location}; got {value!r}.')
    if metric == 'weight' and not 0 <= value <= 100:
        raise ValueError(f'Weight outside 0-100 percent at {location}.')
    if metric == 'duration' and value < 0:
        raise ValueError(f'Negative duration at {location}.')
    return float(value)


def read_breakdowns(path):
    """Exact headers define units; formulas use cached values only, never evaluation."""
    import openpyxl

    values = openpyxl.load_workbook(path, data_only=True, read_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
    records = []
    try:
        for sheet, (asset, kind, headers) in SCHEMAS.items():
            if sheet not in values.sheetnames:
                raise ValueError(f'Missing required sheet: {sheet}.')
            ws, raw = values[sheet], formulas[sheet]
            if tuple(ws.cell(1, c).value for c in range(1, len(headers) + 1)) != headers:
                raise ValueError(f'Unexpected headers in {sheet}; percent and year units must be explicit.')
            seen = set()
            for row in range(2, ws.max_row + 1):
                cells = [ws.cell(row, c) for c in range(1, len(headers) + 1)]
                if all(c.value is None and raw.cell(row, c.column).data_type != 'f' for c in cells):
                    continue
                label = cells[0].value
                if not isinstance(label, str) or not label.strip():
                    raise ValueError(f'Missing category at {sheet}!A{row}.')
                label = label.strip()
                if label.casefold() in seen:
                    raise ValueError(f'Duplicate category {label} in {sheet}.')
                seen.add(label.casefold())
                metrics = ('duration', 'local_yield') if kind == 'currency_analytics' else ('weight',)
                for col, metric in enumerate(metrics, 2):
                    cell = ws.cell(row, col)
                    source = raw.cell(row, col)
                    # Percent-formatted cells store decimals, unlike the export's numeric percent points.
                    if '%' in cell.number_format:
                        raise ValueError(f'Percent-formatted storage at {sheet}!{cell.coordinate}; review units before import.')
                    number = _number(cell.value, f'{sheet}!{cell.coordinate}', metric)
                    records.append({'asset_key': asset, 'breakdown': kind, 'category': label,
                                    'metric': metric, 'source_value': cell.value, 'value': number,
                                    'normalized_value': number if metric == 'duration' else number / 100,
                                    'units': 'years' if metric == 'duration' else 'decimal',
                                    'sheet': sheet, 'label_cell': cells[0].coordinate,
                                    'value_cell': cell.coordinate, 'formula': source.value if source.data_type == 'f' else None,
                                    'status': ('formula_cache_missing' if source.data_type == 'f' and cell.value is None
                                               else 'missing' if math.isnan(number) else 'observed')})
    finally:
        values.close()
        formulas.close()
    return pd.DataFrame(records)


def review_breakdowns(data):
    coverage = []
    for (asset, kind), group in data[data.metric == 'weight'].groupby(['asset_key', 'breakdown']):
        total = group.normalized_value.sum()
        missing = group.normalized_value.isna().sum()
        # Each displayed weight is rounded to two decimal percent points. Never rescale it.
        tolerance = len(group) * .00005 + 1e-12
        unclassified = group.category.str.upper().isin(['NR', 'UNCLASSIFIED'])
        coverage.append({'asset_key': asset, 'breakdown': kind, 'weight_sum_decimal': total,
                         'weight_gap_decimal': 1 - total, 'rounding_tolerance_decimal': tolerance,
                         'total_within_rounding_tolerance': not missing and abs(total - 1) <= tolerance,
                         'missing_weight_rows': int(missing),
                         'unclassified_or_unrated_weight_decimal': group.loc[unclassified, 'normalized_value'].sum(),
                         'full_portfolio_loss_decimal': float('nan')})
    weights = data[(data.breakdown == 'currency') & (data.metric == 'weight')]
    analytics = data[data.breakdown == 'currency_analytics'].pivot(
        index='category', columns='metric', values='normalized_value')
    currencies = weights[['category', 'normalized_value', 'sheet', 'value_cell']].rename(
        columns={'category': 'currency', 'normalized_value': 'weight_decimal'}).merge(
            analytics, left_on='currency', right_index=True, how='left', validate='one_to_one')
    currencies['positive_weight'] = currencies.weight_decimal > 0
    currencies['yield_available_on_positive_weight'] = currencies.positive_weight & currencies.local_yield.notna()
    currencies['duration_available_on_positive_weight'] = currencies.positive_weight & currencies.duration.notna()
    return pd.DataFrame(coverage), currencies


def run_workbook_inputs(output, workbook_path, asof=None, rating_method=None, progress=print):
    path = Path(workbook_path).resolve()
    source_hash = sha256(path)
    date = None
    if asof:
        date = pd.Timestamp(asof)
        if pd.isna(date):
            raise ValueError('Invalid workbook as-of date.')
        date = date.date().isoformat()
    data = read_breakdowns(path)
    if sha256(path) != source_hash:
        raise ValueError('Workbook changed during import; retry with a stable source file.')
    coverage, currencies = review_breakdowns(data)
    folder = save_bundle(data, coverage, output, {
        'kind': 'workbook-input-review', 'source_file': str(path), 'source_sha256': source_hash,
        'asof_date': date, 'asof_source': 'user_supplied' if date else 'not_in_workbook',
        'rating_method': rating_method or 'unconfirmed', 'point_in_time_eligible': False,
        'status': 'research_inputs_only', 'code_sha256': sha256(__file__),
        'weights_rescaled': False, 'joint_sector_rating_distribution_available': False,
        'duration_definition_verified': False, 'source_workbook_modified': False,
        'note': 'Date alone does not approve benchmark definitions, rounding residuals or rating methodology.'})
    currencies.to_csv(folder / 'currency_inputs.csv', index=False)
    for kind in ('sector', 'rating'):
        data[data.breakdown == kind].to_csv(folder / f'{kind}_weights.csv', index=False)
    # These are review templates, not directly executable global-return inputs.
    cash = currencies.loc[currencies.positive_weight, ['currency', 'weight_decimal']].copy()
    for column in ['asof_date', 'usd_cash_decimal', 'foreign_cash_decimal', 'source', 'proxy_note']:
        cash[column] = date if column == 'asof_date' else None
    cash.to_csv(folder / 'currency_cash_review_template.csv', index=False)
    known_yield = currencies.loc[currencies.yield_available_on_positive_weight, 'weight_decimal'].sum()
    known_duration = currencies.loc[currencies.duration_available_on_positive_weight, 'weight_decimal'].sum()
    report = ['# Bloomberg Workbook Input Review', '',
              f'Source: `{path.name}`. SHA-256: `{source_hash}`.', '',
              f'As-of date: {date or "not supplied; file creation time is not an observation date"}. '
              f'Rating methodology: {rating_method or "unconfirmed"}.', '',
              'Source workbook preserved. Percentage-point weights and yields are divided by 100. '
              'Missing markers remain missing, including unavailable formula caches. No weights are rescaled.', '',
              '## Weight Coverage', '', markdown_table(coverage), '',
              'Sector and rating tables are separate marginal distributions, not a sector-by-rating matrix. '
              'NR and Unclassified are not zero-default categories. No corporate PD table or full-index loss is applied.', '',
              '## Foreign Analytics', '',
              f'Positive-weight yield coverage: {known_yield:.2%} of the index. '
              f'Positive-weight duration coverage: {known_duration:.2%} of the index.', '',
              'Numeric yields on zero-weight currencies do not provide usable portfolio yield coverage. '
              'Currency analytics are joined by label, not row order. Duration labels do not establish option-adjusted '
              'rate duration or spread duration. Unclassified currency and displayed rounding gaps remain explicit.', '',
              markdown_table(currencies), '',
              '## Resource-Efficient Global Approach', '',
              'Use the separately verified aggregate index yield as a starting underlying-bond carry proxy. '
              'Per-currency bond yield histories are not required to calculate a cash-differential hedge overlay. '
              'Obtain matched beginning-period USD and foreign cash rates or reviewed forward-implied hedge costs. '
              'The template lists positive-weight exposures, including the unresolved Unclassified exposure.', '',
              '$$', r'y_t^{USD,h,carry}\approx y_t^{aggregate,local}+r_t^{USD}-\sum_c w_{c,t}r_t^c', '$$', '',
              'This is an annualized carry approximation, not total return or an exact hedged index yield. '
              'Do not add it to a return that already includes USD hedging. Currency weights must be reconciled '
              'before calculation. Foreign yield-normalization paths, spread changes, losses and price effects remain '
              'separate inputs. Hedging does not convert foreign duration into US Treasury duration.', '',
              'For long-run foreign rate normalization, use reviewed currency cash anchors plus term/credit premiums '
              'or shorter overlapping government-rate histories with explicit coverage. Do not manufacture a '
              '30-year local-yield history. This import changes no return assumptions and generates no forecast.', '',
              '[JPM methodology, page 31](https://am.jpmorgan.com/content/dam/jpm-am-aem/global/en/insights/portfolio-insights/ltcma/noindex/ltcma-methodology-handbook.pdf)', '']
    (folder / 'run_summary.md').write_text('\n'.join(report), encoding='utf-8')
    manifest = folder / 'manifest.json'
    info = json.loads(manifest.read_text())
    info['output_sha256'] = {p.name: sha256(p) for p in folder.iterdir() if p.name != 'manifest.json'}
    manifest.write_text(json.dumps(info, indent=2), encoding='utf-8')
    if progress:
        progress(f'Saved workbook input review: {folder}')
        progress(f'Local yield coverage: {known_yield:.2%}; duration coverage: {known_duration:.2%}. No forecasts changed.')
    return folder
