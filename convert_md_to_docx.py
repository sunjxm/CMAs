"""Convert CMA Markdown documents into formatted, editable Word review copies."""

import argparse
import glob
import json
from pathlib import Path
import re
import sys
from urllib.parse import urlparse

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parent


def parse_markdown(text):
    try:
        from markdown_it import MarkdownIt
        from mdit_py_plugins.dollarmath import dollarmath_plugin
    except ImportError as exc:
        raise RuntimeError('Install conversion dependencies: python -m pip install -e ".[documents]"') from exc
    parser = MarkdownIt("commonmark", {"html": False}).enable("table").enable("strikethrough")
    parser.use(dollarmath_plugin, allow_labels=False, allow_space=False, allow_digits=True)
    return [token.as_dict() for token in parser.parse(text)]


def add_equation(paragraph, latex, display=False):
    """Convert math through MathML to native, editable Office Math (OMML)."""
    try:
        from latex2mathml.converter import convert as latex_to_mathml
        from mathml2omml import convert as mathml_to_omml
    except ImportError as exc:
        raise RuntimeError('Install equation dependencies: python -m pip install -e ".[documents]"') from exc
    try:
        mathml = latex_to_mathml(latex.strip())
        equation = parse_xml(f"<root {nsdecls('m', 'w')}>{mathml_to_omml(mathml)}</root>")[0]
    except Exception as exc:
        raise ValueError(f"Cannot convert LaTeX equation: {latex!r}: {exc}") from exc
    for run in equation.iter(qn("m:r")):
        properties = OxmlElement("w:rPr")
        font = OxmlElement("w:rFonts")
        font.set(qn("w:ascii"), "Cambria Math")
        font.set(qn("w:hAnsi"), "Cambria Math")
        properties.append(font)
        run.insert(1 if len(run) and run[0].tag == qn("m:rPr") else 0, properties)
    if display:
        container = OxmlElement("m:oMathPara")
        properties = OxmlElement("m:oMathParaPr")
        alignment = OxmlElement("m:jc")
        alignment.set(qn("m:val"), "center")
        properties.append(alignment)
        container.extend([properties, equation])
        paragraph._p.append(container)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(8)
    else:
        paragraph._p.append(equation)


def set_font(run, name=None, size=None, bold=None, italic=None):
    run.bold, run.italic = bold, italic
    if name:
        run.font.name = name
    if size:
        run.font.size = Pt(size)


def inline_text(tokens):
    return "".join(token.get("content", "") for token in tokens
                   if token["type"] in {"text", "code_inline", "math_inline"})


def add_inline(paragraph, tokens, font_size=None):
    bold = italic = strike = 0
    hyperlink = None
    for token in tokens:
        kind = token["type"]
        if kind in {"strong_open", "em_open", "s_open", "strong_close", "em_close", "s_close"}:
            delta = 1 if kind.endswith("open") else -1
            if kind.startswith("strong"):
                bold += delta
            elif kind.startswith("em"):
                italic += delta
            else:
                strike += delta
        elif kind == "link_open":
            href = dict(token.get("attrs") or [])["href"]
            if href.startswith("#"):
                raise ValueError("Internal Markdown anchor links are not supported; use a source URL instead.")
            if urlparse(href).scheme not in {"https", "http", "mailto"}:
                raise ValueError(f"Unsupported link target: {href}")
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
            paragraph._p.append(hyperlink)
        elif kind == "link_close":
            hyperlink = None
        elif kind == "math_inline":
            if hyperlink is not None:
                raise ValueError("Equations inside hyperlinks are not supported.")
            add_equation(paragraph, token["content"])
        elif kind in {"text", "code_inline", "softbreak", "hardbreak"}:
            run = paragraph.add_run()
            if kind in {"softbreak", "hardbreak"}:
                run.add_break() if kind == "hardbreak" else run.add_text(" ")
            else:
                run.text = token["content"]
            set_font(run, "Consolas" if kind == "code_inline" else None,
                     (font_size or 10) if kind == "code_inline" else font_size,
                     True if bold else None, True if italic else None)
            run.font.strike = bool(strike)
            if hyperlink is not None:
                run.font.color.rgb = RGBColor.from_string("24567A")
                run.font.underline = True
                paragraph._p.remove(run._r)
                hyperlink.append(run._r)
        else:
            raise ValueError(f"Unsupported inline Markdown element: {kind}. No content was silently dropped.")


def list_numbering(document, ordered, start):
    numbering = document.part.numbering_part.element
    existing_abstract = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    existing_numbers = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    abstract_id, num_id = max(existing_abstract, default=-1) + 1, max(existing_numbers, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    identity = OxmlElement("w:nsid")
    identity.set(qn("w:val"), f"{0xA0000000 + abstract_id:08X}")
    abstract.append(identity)
    layout = OxmlElement("w:multiLevelType")
    layout.set(qn("w:val"), "singleLevel")
    abstract.append(layout)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in [("start", str(start)), ("numFmt", "decimal" if ordered else "bullet"),
                       ("lvlText", "%1." if ordered else chr(0x2022)), ("lvlJc", "left")]:
        element = OxmlElement("w:" + tag)
        element.set(qn("w:val"), value)
        level.append(element)
    abstract.append(level)
    # OOXML requires all abstract definitions before concrete numbering instances.
    first_number = numbering.find(qn("w:num"))
    numbering.insert(list(numbering).index(first_number) if first_number is not None else len(numbering), abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    number.append(reference)
    numbering.append(number)
    return num_id


def apply_list(paragraph, num_id, depth):
    properties = paragraph._p.get_or_add_pPr()
    num = OxmlElement("w:numPr")
    level, reference = OxmlElement("w:ilvl"), OxmlElement("w:numId")
    level.set(qn("w:val"), "0")
    reference.set(qn("w:val"), str(num_id))
    num.extend([level, reference])
    properties.append(num)
    paragraph.paragraph_format.left_indent = Inches(.22 * depth)
    paragraph.paragraph_format.first_line_indent = Inches(-.17)


def configure_document(document, landscape=False):
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    section.footer_distance = Inches(.3)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6"):
        style = document.styles[name]
        for borders in style._element.findall(".//" + qn("w:pBdr")):
            borders.getparent().remove(borders)
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt({"Normal": 10.5, "Title": 24, "Heading 1": 15,
                             "Heading 2": 12, "Heading 3": 11}.get(name, 11))
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.10
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.keep_with_next = True
        if name == "Title":
            style.paragraph_format.space_after = Pt(14)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Page ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""


def extract_table(tokens, start):
    rows, row, index = [], [], start + 1
    while index < len(tokens) and tokens[index]["type"] != "table_close":
        token = tokens[index]
        if token["type"] == "tr_open":
            row = []
        elif token["type"] == "inline":
            row.append(token.get("children") or [])
        elif token["type"] == "tr_close":
            rows.append(row)
        index += 1
    if not rows or any(len(row) != len(rows[0]) for row in rows):
        raise ValueError("Malformed Markdown table.")
    return rows, index


def add_table(document, rows):
    columns = len(rows[0])
    section = document.sections[-1]
    available = (section.page_width - section.left_margin - section.right_margin) / 914400
    # Allocate more room to explanatory columns than compact dates or numeric fields.
    weights = []
    for column in range(columns):
        values = [inline_text(row[column]) for row in rows]
        words = [len(word) for value in values for word in value.replace("_", " ").split()]
        average = sum(min(len(value), 80) for value in values) / len(values)
        weights.append(max(10, min(45, average * .6 + max(words, default=1))))
    minimum = .65
    remaining = available - minimum * columns
    widths = [minimum + remaining * weight / sum(weights) for weight in weights]
    table = document.add_table(rows=len(rows), cols=columns)
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        for name, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn("w:" + name), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for row_index, (row, contents) in enumerate(zip(table.rows, rows)):
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            properties.append(OxmlElement("w:tblHeader"))
        for column_index, (cell, content, width) in enumerate(zip(row.cells, contents, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + edge)
                margin.set(qn("w:w"), "85")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell_properties.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E7EBEF" if row_index == 0 else "FFFFFF")
            cell_properties.append(shading)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            paragraph.paragraph_format.keep_with_next = row_index == 0
            text = inline_text(content)
            if row_index > 0 and (re.fullmatch(r"[\d.,%+-]+", text) or re.fullmatch(r"\d{4}-\d{2}-\d{2}", text)):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Display schema headers with spaces so long technical names can wrap.
            content = [dict(token, content=token["content"].replace("_", " "))
                       if row_index == 0 and token["type"] == "text" else token for token in content]
            add_inline(paragraph, content, font_size=9)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build_document(tokens, title="", orientation="auto"):
    maximum_columns = max((sum(t["type"] == "th_open" for t in tokens[i + 1:j])
                           for i, token in enumerate(tokens) if token["type"] == "table_open"
                           for j in [next(k for k in range(i + 1, len(tokens)) if tokens[k]["type"] == "thead_close")]), default=0)
    if maximum_columns > 10:
        raise ValueError("Tables with more than 10 columns require a dedicated layout.")
    landscape = orientation == "landscape" or (orientation == "auto" and maximum_columns >= 6)
    document = Document()
    configure_document(document, landscape)
    document.core_properties.title = title
    lists, index, quote_depth = [], 0, 0
    title_added = False
    while index < len(tokens):
        token = tokens[index]
        kind = token["type"]
        if kind == "heading_open":
            level = int(token["tag"][1:])
            inline = tokens[index + 1]
            is_title = level == 1 and not title_added
            paragraph = document.add_paragraph(style="Title" if is_title else f"Heading {max(1, level - 1)}")
            add_inline(paragraph, inline.get("children") or [])
            title_added = title_added or is_title
            index += 2
        elif kind == "paragraph_open":
            paragraph = document.add_paragraph()
            if quote_depth:
                paragraph.paragraph_format.left_indent = Inches(.25 * quote_depth)
            if lists:
                item = lists[-1]
                if item["pending"]:
                    apply_list(paragraph, item["num_id"], len(lists))
                    item["pending"] = False
                else:
                    paragraph.paragraph_format.left_indent = Inches(.22 * len(lists))
            add_inline(paragraph, tokens[index + 1].get("children") or [])
            index += 2
        elif kind in {"bullet_list_open", "ordered_list_open"}:
            start = int(dict(token.get("attrs") or []).get("start", 1))
            lists.append({"num_id": list_numbering(document, kind == "ordered_list_open", start), "pending": False})
        elif kind == "list_item_open":
            lists[-1]["pending"] = True
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif kind == "table_open":
            rows, index = extract_table(tokens, index)
            add_table(document, rows)
        elif kind == "math_block":
            add_equation(document.add_paragraph(), token["content"], display=True)
        elif kind in {"fence", "code_block"}:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(.12)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(3)
            set_font(paragraph.add_run(token["content"].rstrip("\n")), "Consolas", 9)
        elif kind == "blockquote_open":
            quote_depth += 1
        elif kind == "blockquote_close":
            quote_depth -= 1
        elif kind == "hr":
            document.add_paragraph()
        elif kind not in {"list_item_close", "heading_close", "paragraph_close"}:
            raise ValueError(f"Unsupported Markdown block: {kind}. No content was silently dropped.")
        index += 1
    return document


def convert_file(source, destination, overwrite=False, orientation="auto"):
    source, destination = Path(source), Path(destination)
    if source.suffix.lower() != ".md" or not source.is_file():
        raise ValueError(f"Markdown file not found: {source}")
    if destination.exists() and not overwrite:
        raise ValueError(f"Output already exists: {destination}. Use --overwrite to refresh it.")
    tokens = parse_markdown(source.read_text(encoding="utf-8-sig"))
    document = build_document(tokens, title=source.stem, orientation=orientation)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


def default_sources(project):
    sources = sorted((project / "docs").glob("*.md"))
    reviews = []
    for manifest_path in (project / "data").glob("*/manifest.json"):
        metadata = json.loads(manifest_path.read_text(encoding="utf-8"))
        if metadata.get("kind") == "anchor-review":
            reviews.append((metadata["created_at_utc"], manifest_path.parent / "run_summary.md"))
    if reviews:
        sources.append(max(reviews, key=lambda item: item[0])[1])
    return sources


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("files", nargs="*", help="Markdown files, directories, or quoted glob patterns.")
    parser.add_argument("--output", type=Path, default=PROJECT_DIR / "docs" / "word")
    parser.add_argument("--all", action="store_true", help="Also include project-root Markdown files.")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--orientation", choices=["auto", "portrait", "landscape"], default="auto")
    args = parser.parse_args(argv)
    sources = []
    for value in args.files:
        matches = sorted(glob.glob(value))
        if not matches:
            raise ValueError(f"Input path/pattern does not match: {value}")
        for match in matches:
            path = Path(match)
            sources.extend(sorted(path.glob("*.md")) if path.is_dir() else [path])
    if not args.files:
        sources = default_sources(PROJECT_DIR)
    if args.all:
        sources.extend(sorted(PROJECT_DIR.glob("*.md")))
    sources = list(dict.fromkeys(path.resolve() for path in sources))
    if not sources:
        raise ValueError("No Markdown files selected.")
    targets = [args.output / (source.stem + ".docx") for source in sources]
    if len({str(path).lower() for path in targets}) != len(targets):
        raise ValueError("Selected files have duplicate output names. Convert them into separate output folders.")
    for source, destination in zip(sources, targets):
        result = convert_file(source, destination, args.overwrite, args.orientation)
        print(f"Created: {result.resolve()}", flush=True)


if __name__ == "__main__":
    try:
        main()
    except (ValueError, RuntimeError, OSError) as exc:
        print(f"Word conversion failed: {exc}", file=sys.stderr)
        sys.exit(1)
