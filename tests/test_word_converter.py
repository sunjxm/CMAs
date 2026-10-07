from pathlib import Path
import tempfile
import unittest
import zipfile

from docx.oxml.ns import qn
from convert_md_to_docx import build_document, convert_file, parse_markdown


class WordConverterTests(unittest.TestCase):
    def test_display_math_creates_editable_fraction_and_superscript(self):
        document = build_document(parse_markdown(r"$$" + "\n" + r"y_m(h) = a_m + 2^{-\frac{h}{H_m}}" + "\n$$"))
        xml = document.paragraphs[0]._p.xml
        self.assertIn("m:oMathPara", xml)
        self.assertIn("m:f>", xml)
        self.assertIn("m:sSup", xml)
        self.assertIn("Cambria Math", xml)
        self.assertNotIn(r"\frac", xml)

    def test_inline_math_preserves_surrounding_text_and_table_equation(self):
        document = build_document(parse_markdown("Rate $r^{*}$ is real.\n\n| Symbol |\n|---|\n| $x_1$ |"))
        self.assertEqual(document.paragraphs[0].text, "Rate  is real.")
        self.assertIn("m:oMath", document.paragraphs[0]._p.xml)
        self.assertNotIn("m:oMathPara", document.paragraphs[0]._p.xml)
        self.assertIn("m:sSub", document.tables[0].cell(1, 0)._tc.xml)

    def test_currency_and_code_are_not_math(self):
        source = r"Costs \$5 and \$10; use `$r^{*}$`." + "\n\n```text\n$$x$$\n```"
        document = build_document(parse_markdown(source))
        self.assertEqual(document.paragraphs[0].text, "Costs $5 and $10; use $r^{*}$.")
        self.assertEqual(document.paragraphs[-1].text, "$$x$$")
        self.assertNotIn("m:oMath", document._element.xml)

    def test_math_survives_docx_serialization(self):
        document = build_document(parse_markdown("$$\n" + r"\widehat{s}_m = \frac{1}{N_m} \sum_{t \in \mathcal{T}_m} s_{m,t}" + "\n$$"))
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as directory:
            output = Path(directory) / "equation.docx"
            document.save(output)
            with zipfile.ZipFile(output) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("m:oMathPara", xml)
            self.assertIn("m:nary", xml)
            self.assertNotIn(r"\sum", xml)

    def test_styles_emphasis_and_clickable_links(self):
        document = build_document(parse_markdown("# Title\n\n## Heading\n\n**Bold** and *italic* with `code` and [source](https://example.com)."))
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertNotIn("pBdr", document.styles["Title"]._element.xml)
        self.assertEqual(document.paragraphs[1].style.name, "Heading 1")
        paragraph = document.paragraphs[2]
        self.assertTrue(any(r.bold and r.text == "Bold" for r in paragraph.runs))
        self.assertTrue(any(r.italic and r.text == "italic" for r in paragraph.runs))
        self.assertTrue(any(r.font.name == "Consolas" for r in paragraph.runs))
        self.assertIn("w:hyperlink", paragraph._p.xml)
        self.assertTrue(any(r.target_ref == "https://example.com" for r in document.part.rels.values()))

    def test_table_values_header_repeat_borders_and_width(self):
        document = build_document(parse_markdown("# Table\n\n| key | yield_percent |\n|---|---|\n| cash | 3.15% |"))
        table = document.tables[0]
        self.assertEqual(table.cell(0, 1).text, "yield percent")
        self.assertEqual(table.cell(1, 1).text, "3.15%")
        self.assertIn("tblHeader", table.rows[0]._tr.xml)
        self.assertIn('w:color="D9D9D9"', table._tbl.xml)
        self.assertFalse(table.autofit)

    def test_real_list_numbering_and_literal_code_block(self):
        source = "# Lists\n\n3. First\n4. Second\n\n- Bullet\n\n```python\nx = 2 ** -3\n```"
        document = build_document(parse_markdown(source))
        self.assertIn("numPr", document.paragraphs[1]._p.xml)
        self.assertIn('w:val="3"', document.part.numbering_part.element.xml)
        children = list(document.part.numbering_part.element)
        first_number = next(i for i, child in enumerate(children) if child.tag == qn("w:num"))
        self.assertTrue(all(child.tag != qn("w:abstractNum") for child in children[first_number:]))
        self.assertEqual(document.paragraphs[-1].text, "x = 2 ** -3")

    def test_wide_table_landscape_and_explicit_portrait(self):
        text = "| A | B | C | D | E | F |\n|---|---|---|---|---|---|\n|1|2|3|4|5|6|"
        tokens = parse_markdown(text)
        self.assertGreater(build_document(tokens).sections[0].page_width, build_document(tokens).sections[0].page_height)
        document = build_document(tokens, orientation="portrait")
        self.assertLess(document.sections[0].page_width, document.sections[0].page_height)

    def test_file_round_trip_preserves_source_and_guards_overwrite(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as directory:
            source, output = Path(directory) / "sample.md", Path(directory) / "sample.docx"
            content = "# Source\n\nSource is unchanged."
            source.write_text(content, encoding="utf-8")
            convert_file(source, output)
            with zipfile.ZipFile(output) as archive:
                self.assertIn("word/document.xml", archive.namelist())
            self.assertEqual(source.read_text(encoding="utf-8"), content)
            with self.assertRaisesRegex(ValueError, "already exists"):
                convert_file(source, output)
            convert_file(source, output, overwrite=True)

    def test_unsupported_images_fail_instead_of_silently_dropping_content(self):
        with self.assertRaisesRegex(ValueError, "image"):
            build_document(parse_markdown("![Chart](chart.png)"))
